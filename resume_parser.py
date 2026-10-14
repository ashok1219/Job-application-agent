from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

MAX_RESUME_BYTES = 10 * 1024 * 1024


def extract_resume_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".pdf", ".docx", ".txt"}:
        raise ValueError("Upload a PDF, DOCX, or TXT resume.")
    if len(content) > MAX_RESUME_BYTES:
        raise ValueError("Resume file must be 10 MB or smaller.")

    if suffix == ".pdf":
        try:
            text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(content)).pages)
        except PdfReadError as exc:
            raise ValueError("Could not read this PDF. Check that the file is valid and not password-protected.") from exc
    elif suffix == ".docx":
        try:
            document = Document(BytesIO(content))
        except (BadZipFile, PackageNotFoundError) as exc:
            raise ValueError("Could not read this DOCX. Check that the file is valid and not password-protected.") from exc
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        tables = [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
        text = "\n".join([*paragraphs, *tables])
    else:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("Could not read the TXT file as UTF-8 text.") from exc

    if not text.strip():
        raise ValueError("No selectable text was found. Scanned/image-only PDFs are not supported.")
    return text
