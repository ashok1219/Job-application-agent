import unittest
from io import BytesIO

from docx import Document
from pypdf import PdfWriter

from resume_parser import extract_resume_text


class ResumeParserTests(unittest.TestCase):
    def test_extracts_utf8_text(self):
        text = extract_resume_text("resume.txt", "Python and SQL".encode())

        self.assertEqual("Python and SQL", text)

    def test_extracts_docx_paragraphs_and_tables(self):
        document = Document()
        document.add_paragraph("QA Engineer")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Selenium"
        content = BytesIO()
        document.save(content)

        text = extract_resume_text("resume.docx", content.getvalue())

        self.assertIn("QA Engineer", text)
        self.assertIn("Selenium", text)

    def test_rejects_pdf_without_extractable_text(self):
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        content = BytesIO()
        writer.write(content)

        with self.assertRaisesRegex(ValueError, "No selectable text"):
            extract_resume_text("resume.pdf", content.getvalue())

    def test_rejects_unsupported_and_empty_files(self):
        with self.assertRaisesRegex(ValueError, "PDF, DOCX, or TXT"):
            extract_resume_text("resume.rtf", b"content")
        with self.assertRaisesRegex(ValueError, "No selectable text"):
            extract_resume_text("resume.txt", b"  ")


if __name__ == "__main__":
    unittest.main()
